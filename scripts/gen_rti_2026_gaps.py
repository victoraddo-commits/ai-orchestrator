#!/usr/bin/env python3
"""Generate the 2026-gap RTI pack: Standard Form (Appendix A) + cover email.

Closes the newest Juris Kai legal-brain corpus gaps by requesting the *official
electronic text* of instruments the corpus does not yet hold. Ghana enactments
(Acts and legislative instruments) are excluded from copyright under s.8(1)(a)
of the Copyright Act, 2005 (Act 690); the request seeks the official text
itself, not any editorial matter.

Fields mirror docs/requests/rti_application_standard_form_2026-09-24.md
(RTI Commission Information Manual 2022, Appendix A). Commercial framing mirrors
the existing Juris Kai access letters. Renders one PDF + one DOCX per request
into docs/requests/ via fpdf2 + python-docx (same stack as core/juris_kai/reports.py).

Run:  cd /opt/ai-orchestrator && PYTHONPATH=. .venv/bin/python scripts/gen_rti_2026_gaps.py
"""
from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / "docs" / "requests"
PREPARED = date.today().isoformat()

_FONT_DIRS = (
    "/usr/share/fonts/truetype/dejavu",
    "/usr/share/fonts/dejavu",
    "/usr/share/fonts/TTF",
)

COMMERCIAL_PURPOSE = (
    "Juris Kai is a paid (commercial) legal-research and legal-information "
    "service for Ghanaian law, operated by [ORGANISATION]. The requested text is "
    "the official legal instrument itself; enactments and legislative instruments "
    "are excluded from copyright under s.8(1)(a) of the Copyright Act, 2005 "
    "(Act 690). We seek the official electronic text to answer questions "
    "accurately from the official source, will attribute the institution as the "
    "source, will not misrepresent the text or present it as legal advice, and "
    "will use it subject to any conditions the institution specifies."
)

REQUESTS = [
    {
        "key": "Road_Traffic_Regulations_LI2519_MinTransport",
        "short": "Road Traffic Regulations, 2026 (L.I. 2519) + predecessor L.I. 2180",
        "institution": "Ministry of Transport",
        "cc": "Driver and Vehicle Licensing Authority (DVLA); National Road Safety Authority (NRSA)",
        "class": "Subsidiary legislation (Legislative Instrument)",
        "instruments": [
            "Road Traffic Regulations, 2026 (L.I. 2519) — the current instrument;",
            "Road Traffic Regulations, 2012 (L.I. 2180) — the predecessor instrument, "
            "together with any amendments made to it before 2026.",
        ],
        "cover_dates": "2012 to date (L.I. 2180 and its amendments; L.I. 2519, 2026)",
        "subject": "RTI application (Act 989) — official text of the Road Traffic "
                   "Regulations, 2026 (L.I. 2519) and L.I. 2180",
        "gap": "Legal-brain transport gaps #1/#3 — 'what % tint is allowed on cars?' / "
               "'Traffic & vehicle tint' (missing L.I. 2519 and L.I. 2180).",
    },
    {
        "key": "Tax_Levy_Acts_2026_MinFinance_GRA",
        "short": "2026 tax and levy Acts (VAT, Income Tax, Customs, Excise, Energy "
                 "Sector Levies, Growth & Sustainability Levy)",
        "institution": "Ministry of Finance",
        "cc": "Ghana Revenue Authority (GRA)",
        "class": "Primary legislation (Acts of Parliament)",
        "instruments": [
            "Value Added Tax (Amendment) Act, 2026;",
            "Income Tax (Amendment) Act, 2026;",
            "Customs Act, 2026;",
            "Excise Act, 2026;",
            "Energy Sector Levies (Amendment) Act, 2026;",
            "Growth and Sustainability Levy (Amendment) Act, 2026.",
        ],
        "cover_dates": "1 January 2026 to date",
        "subject": "RTI application (Act 989) — official text of the 2026 tax and "
                   "levy Acts",
        "gap": "Corpus holds tax/levy Bills and pre-2026 amending Acts but none of the "
               "2026 tax/levy Acts (VAT, Income Tax, Customs, Excise, Energy Sector "
               "Levies, Growth & Sustainability Levy).",
    },
    {
        "key": "Value_for_Money_Office_Act_PPA_MoF",
        "short": "Value for Money Office Act, 2026",
        "institution": "Public Procurement Authority",
        "cc": "Ministry of Finance",
        "class": "Primary legislation (Act of Parliament)",
        "instruments": ["Value for Money Office Act, 2026."],
        "cover_dates": "2026 to date",
        "subject": "RTI application (Act 989) — official text of the Value for "
                   "Money Office Act, 2026",
        "gap": "No 'Value for Money Office Act' in the public-finance corpus.",
    },
    {
        "key": "Legal_Education_Reform_Act_2025_MinEducation",
        "short": "Legal Education Reform Act, 2025",
        "institution": "Ministry of Education",
        "cc": "Ghana School of Law",
        "class": "Primary legislation (Act of Parliament)",
        "instruments": ["Legal Education Reform Act, 2025."],
        "cover_dates": "2025 to date",
        "subject": "RTI application (Act 989) — official text of the Legal "
                   "Education Reform Act, 2025",
        "gap": "No 'Legal Education Reform Act' in the legal-education corpus.",
    },
    {
        "key": "National_Defence_University_Act_2026_MinDefence",
        "short": "National Defence University Act, 2026",
        "institution": "Ministry of Defence",
        "cc": "",
        "class": "Primary legislation (Act of Parliament)",
        "instruments": ["National Defence University Act, 2026."],
        "cover_dates": "2026 to date",
        "subject": "RTI application (Act 989) — official text of the National "
                   "Defence University Act, 2026",
        "gap": "No 'National Defence University Act' in the corpus.",
    },
]


# ── shared text ─────────────────────────────────────────────────────────────

def _form_lines(req: dict) -> list[tuple[str, str]]:
    """The completed Standard Form (Appendix A) as (kind, text) blocks."""
    inst = "\n".join("        • " + x for x in req["instruments"])
    cc_line = ("Cc: " + req["cc"]) if req["cc"] else ""
    return [
        ("h1", "APPLICATION FOR ACCESS TO INFORMATION"),
        ("sub", "Right to Information Act, 2019 (Act 989) — Standard Form (Appendix A)"),
        ("small", "[Reference No.: ……………………………………]   (assigned by the institution)"),
        ("kv", "1.  Name of Applicant:  [FULL NAME]"),
        ("kv", "2.  Date:  [DATE]"),
        ("kv", f"3.  Public Institution:  {req['institution']}"),
        ("kv", f"    {cc_line}" if cc_line else "    "),
        ("kv", "4.  Date of Birth:  [DD/MM/YYYY]"),
        ("kv", "5.  Type of Applicant:  [ ] Individual      [X] Organization/Institution"),
        ("kv", "6.  Tax Identification Number:  [TIN, if applicable]"),
        ("kv", "7.  If Represented, Name of Person Being Represented:  —"),
        ("kv", "7(a). Capacity of Representative:  —"),
        ("kv", "8.  Type of Identification:  [X] National ID Card  [ ] Passport  "
               "[ ] Voter's ID  [ ] Driver's Licence"),
        ("kv", "8(a). ID No.:  [ID NUMBER]"),
        ("blank", ""),
        ("kv", "9.  Description of the Information being sought"),
        ("kv", "    (specify the type and class of information including cover dates):"),
        ("kv", f"    Class:  {req['class']}."),
        ("kv", "    Type:  Official full text in electronic (machine-readable) form "
               "— PDF and, where available, the gazetted text — together with "
               "commencement and amendment metadata (date made, date of publication "
               "in the Gazette, and any amendments)."),
        ("kv", "    Instruments:"),
        ("pre", inst),
        ("kv", f"    Cover dates:  {req['cover_dates']}."),
        ("kv", "    Purpose:  " + COMMERCIAL_PURPOSE),
        ("blank", ""),
        ("kv", "10. Manner of Access:"),
        ("kv", "        [ ] Inspection of Information   [X] Copy of Information"),
        ("kv", "        [ ] Viewing / Listen            [ ] Written Transcript"),
        ("kv", "        [ ] Translated (specify language):  —"),
        ("kv", "10(a). Form of Access:"),
        ("kv", "        [ ] Hard copy   [X] Electronic copy (transmitted by email)   [ ] Braille"),
        ("blank", ""),
        ("kv", "11. Contact Details:"),
        ("kv", "        Email Address:  [EMAIL]"),
        ("kv", "        Postal Address:  [POSTAL ADDRESS]"),
        ("kv", "        Tel:  [PHONE]"),
        ("kv", "12. Applicant's signature/thumbprint:  ……………………………………"),
        ("blank", ""),
        ("kv", "13. Signature of Witness (where applicable)"),
        ("small", "\"This request was read to the applicant in the language the applicant "
                  "understands and the applicant appeared to have understood the content "
                  "of the request.\""),
        ("kv", "    Witness:  ……………………………………   Date:  ……………………"),
        ("blank", ""),
        ("small", "Statutory note: a decision is due within 14 days (s.23); transfer "
                  "within 10 days (s.20); an extension is limited to 7 days. The "
                  "applicable fee (s.24) will be paid on request. The official text "
                  "requested is not copyrightable (Copyright Act, 2005 (Act 690), "
                  "s.8(1)(a))."),
    ]


def _email_lines(req: dict) -> list[tuple[str, str]]:
    to = f"The RTI Officer, {req['institution']}"
    body = [
        ("kv", f"To:  {to}"),
        ("kv", f"Cc:  {req['cc']}" if req["cc"] else "Cc:  —"),
        ("kv", f"Subject:  {req['subject']}"),
        ("blank", ""),
        ("p", f"Dear RTI Officer,"),
        ("p", "Please find below a completed Standard RTI Application Form "
              "(Appendix A) under the Right to Information Act, 2019 (Act 989), "
              f"addressed to the {req['institution']}."),
        ("p", "In summary, the request seeks the official text of the following "
              "instrument(s) in electronic copy, transmitted by email, for a paid "
              "(commercial) legal-research and legal-information service:"),
        ("pre", "\n".join("    • " + x for x in req["instruments"])),
        ("p", "Enactments and legislative instruments are excluded from copyright "
              "under s.8(1)(a) of the Copyright Act, 2005 (Act 690). The request "
              "concerns the official text of the instrument itself, not any "
              "editorial matter. We will attribute the institution as the source, "
              "will not republish the text in bulk or misrepresent it, and will "
              "comply with any conditions you specify."),
        ("p", "I confirm readiness to pay the applicable fee (s.24). Kindly "
              "acknowledge receipt and provide a written decision within the "
              "14-day period in s.23. If any part is exempt, please sever and "
              "provide the non-exempt part with the ground of refusal, and advise "
              "the internal review / Right to Information Commission appeal route."),
        ("p", "Yours faithfully,"),
        ("kv", "[FULL NAME]"),
        ("kv", "[TITLE], [ORGANISATION]"),
        ("kv", "[EMAIL] | [PHONE] | [POSTAL ADDRESS]"),
        ("blank", ""),
        ("small", "Attachments: (1) completed Standard RTI Application Form; "
                  "(2) copy of valid ID."),
    ]
    return [
        ("h1", "COVER EMAIL (submit with the completed form)"),
        *body,
    ]


def _blocks(req: dict) -> list[tuple[str, str]]:
    return _form_lines(req) + [("blank", ""), ("hr", ""), ("blank", "")] + _email_lines(req)


# ── PDF (fpdf2) ─────────────────────────────────────────────────────────────

def _pdf_font(pdf):
    for directory in _FONT_DIRS:
        regular = Path(directory) / "DejaVuSans.ttf"
        bold = Path(directory) / "DejaVuSans-Bold.ttf"
        if regular.exists():
            try:
                pdf.add_font("DejaVu", "", str(regular))
                if bold.exists():
                    pdf.add_font("DejaVu", "B", str(bold))
                return "DejaVu"
            except Exception:  # noqa: BLE001
                break
    return "Helvetica"


def render_pdf(req: dict) -> bytes:
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    font = _pdf_font(pdf)

    def out(text, size=10, style=""):
        pdf.set_font(font, style, size)
        pdf.multi_cell(0, 5.4, str(text), new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    out("Juris Kai — RTI Application Pack", 15, "B")
    out(f"Prepared {PREPARED} · {req['short']}", 10)
    pdf.ln(1)

    for kind, text in _blocks(req):
        if kind == "h1":
            pdf.ln(1); out(text, 12.5, "B")
        elif kind == "sub":
            out(text, 10.5, "B")
        elif kind == "small":
            out(text, 8.5)
        elif kind == "pre":
            out(text, 9.5)
        elif kind == "hr":
            out("―" * 60, 10)
        elif kind == "blank":
            pdf.ln(1)
        else:
            out(text, 9.6)
    return bytes(pdf.output())


# ── DOCX (python-docx) ──────────────────────────────────────────────────────

def render_docx(req: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading("Juris Kai — RTI Application Pack", level=0)
    doc.add_paragraph(f"Prepared {PREPARED} · {req['short']}")

    for kind, text in _blocks(req):
        if kind == "h1":
            doc.add_heading(text, level=1)
        elif kind == "sub":
            p = doc.add_paragraph(); p.add_run(text).bold = True
        elif kind == "small":
            p = doc.add_paragraph(); r = p.add_run(text); r.font.size = Pt(8.5)
        elif kind == "pre":
            for ln in str(text).splitlines() or [""]:
                p = doc.add_paragraph(); r = p.add_run(ln); r.font.name = "Consolas"
        elif kind == "hr":
            doc.add_paragraph("―" * 60)
        elif kind == "blank":
            doc.add_paragraph("")
        else:
            doc.add_paragraph(text)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_index(reqs) -> str:
    lines = [
        "# RTI pack — 2026 corpus gaps",
        "",
        f"Prepared {PREPARED}. Standard Form (Appendix A) under the Right to "
        "Information Act, 2019 (Act 989) + cover email, one PDF and one DOCX per "
        "request. Enactments are copyright-excluded (Copyright Act, 2005 (Act 690), "
        "s.8(1)(a)); each request seeks the official electronic text. Commercial "
        "framing: Juris Kai is a paid legal-research service.",
        "",
        "| # | Request | Recipient (cc) | Corpus gap closed | PDF | DOCX |",
        "|---|---------|----------------|-------------------|-----|------|",
    ]
    for i, r in enumerate(reqs, 1):
        pdf = f"{r['key']}.pdf"
        docx = f"{r['key']}.docx"
        cc = f" (cc {r['cc']})" if r["cc"] else ""
        gap = r["gap"].replace("|", "/")
        lines.append(f"| {i} | {r['short']} | {r['institution']}{cc} | {gap} | "
                     f"[{pdf}]({pdf}) | [{docx}]({docx}) |")
    lines += [
        "",
        "## Send checklist (per request)",
        "- Fill `[BRACKETS]`: full name, date, DOB, TIN, ID type/number, email, "
        "phone, postal address; sign box 12 (witness only if read orally).",
        "- Download the official blank Standard RTI Application Form (Appendix A) "
        "and transcribe these fields, or submit the electronic copy.",
        "- Attach a copy of one valid ID (Passport / National ID / Voter's ID / "
        "Driver's Licence).",
        "- Tick Manner of Access = Copy of Information; Form = Electronic copy.",
        "- Note the clock: decision within 14 days (s.23); transfer within 10 days "
        "(s.20); extension max 7 days; fee per s.24.",
        "- Verify the current recipient email/postal address before sending.",
        "",
        "Generated by `scripts/gen_rti_2026_gaps.py`.",
        "",
    ]
    return "\n".join(lines)


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    for r in REQUESTS:
        pdf_bytes = render_pdf(r)
        docx_bytes = render_docx(r)
        (OUT / f"{r['key']}.pdf").write_bytes(pdf_bytes)
        (OUT / f"{r['key']}.docx").write_bytes(docx_bytes)
        print(f"{r['key']:<52} pdf={len(pdf_bytes):>7}B docx={len(docx_bytes):>7}B "
              f"magic={pdf_bytes[:4]!r}/{docx_bytes[:2]!r}")
    (OUT / "README_rti_2026_gaps.md").write_text(render_index(REQUESTS), encoding="utf-8")
    print("index -> docs/requests/README_rti_2026_gaps.md")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
